"""DOCX parser: paragraph text via python-docx. Word has no reliable page map."""

from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.exceptions import ParseFailedError
from app.schemas.parsed import Page, ParsedDocument, ParsedDocumentMetadata


def parse_docx(path: Path, *, filename: str) -> ParsedDocument:
    """Parse a DOCX file into a normalized `ParsedDocument`.

    python-docx does not expose print pagination. Paragraphs are concatenated and
    emitted as a single logical page so later chunks still carry `page=1` rather
    than fabricating page numbers.
    """
    try:
        document = Document(str(path))
    except (PackageNotFoundError, ValueError, KeyError, OSError) as exc:
        raise ParseFailedError(
            "Unable to parse DOCX.",
            detail={"filename": filename, "reason": type(exc).__name__},
        ) from exc

    paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
    # Tables are industrial-doc critical; flatten cell text so it is not dropped.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    body = "\n".join(paragraphs)
    pages = [Page(page_number=1, text=body)]
    return ParsedDocument(
        text=body,
        metadata=ParsedDocumentMetadata(
            filename=filename,
            page_count=1,
            source_type="docx",
        ),
        pages=pages,
    )
