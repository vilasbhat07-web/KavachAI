"""PDF/DOCX/TXT/CSV parsers emit ParsedDocument; unsupported types raise 415."""

from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

from app.config import Settings
from app.exceptions import ParseFailedError, UnsupportedMediaTypeError
from app.parsers import parse_document
from app.parsers.csv_parser import parse_csv
from app.parsers.docx_parser import parse_docx
from app.parsers.txt_parser import parse_txt


@pytest.fixture
def settings() -> Settings:
    return Settings(_env_file=None)


def test_txt_parser_preserves_body(tmp_path: Path) -> None:
    path = tmp_path / "note.txt"
    path.write_text("line one\nline two\n", encoding="utf-8")
    parsed = parse_txt(path, filename="note.txt")
    assert parsed.metadata.source_type == "txt"
    assert parsed.metadata.page_count == 1
    assert parsed.pages[0].page_number == 1
    assert "line one" in parsed.text


def test_csv_parser_exposes_columns(tmp_path: Path) -> None:
    path = tmp_path / "rows.csv"
    path.write_text("temp,pressure\n20,1.0\n21,1.1\n", encoding="utf-8")
    parsed = parse_csv(path, filename="rows.csv")
    assert parsed.metadata.source_type == "csv"
    assert parsed.extra is not None
    assert parsed.extra["columns"] == ["temp", "pressure"]
    assert parsed.extra["row_count"] == 2
    assert "20 | 1.0" in parsed.text


def test_docx_parser_includes_tables(tmp_path: Path) -> None:
    path = tmp_path / "spec.docx"
    doc = Document()
    doc.add_paragraph("Header text")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "kPa"
    table.rows[0].cells[1].text = "101"
    doc.save(path)
    parsed = parse_docx(path, filename="spec.docx")
    assert parsed.metadata.source_type == "docx"
    assert "Header text" in parsed.text
    assert "kPa | 101" in parsed.text


def test_pdf_parser_keeps_page_numbers(tmp_path: Path, settings: Settings) -> None:
    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.add_blank_page(width=72, height=72)
    with path.open("wb") as handle:
        writer.write(handle)
    parsed = parse_document(
        path,
        filename="blank.pdf",
        content_type="application/pdf",
        settings=settings,
    )
    assert parsed.metadata.source_type == "pdf"
    assert parsed.metadata.page_count == 2
    assert [p.page_number for p in parsed.pages] == [1, 2]


def test_dispatch_rejects_gif_with_415(tmp_path: Path, settings: Settings) -> None:
    path = tmp_path / "pic.gif"
    path.write_bytes(b"GIF89a")
    with pytest.raises(UnsupportedMediaTypeError) as exc:
        parse_document(path, filename="pic.gif", content_type="image/gif", settings=settings)
    assert exc.value.http_status == 415
    assert exc.value.error_code == "UNSUPPORTED_MEDIA_TYPE"


def test_dispatch_octet_stream_uses_extension(tmp_path: Path, settings: Settings) -> None:
    path = tmp_path / "note.txt"
    path.write_text("ok", encoding="utf-8")
    parsed = parse_document(
        path,
        filename="note.txt",
        content_type="application/octet-stream",
        settings=settings,
    )
    assert parsed.metadata.source_type == "txt"


def test_missing_file_is_parse_failed(tmp_path: Path, settings: Settings) -> None:
    with pytest.raises(ParseFailedError) as exc:
        parse_document(tmp_path / "gone.txt", filename="gone.txt", content_type="text/plain", settings=settings)
    assert exc.value.http_status == 422
